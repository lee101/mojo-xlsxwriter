from __future__ import annotations

import datetime as dt
import math
import zipfile
from io import BytesIO
from xml.etree import ElementTree as ET

import openpyxl
import pytest
import xlsxwriter
from xlsxwriter.utility import (
    xl_col_to_name as upstream_col_name,
    xl_range as upstream_range,
    xl_range_abs as upstream_range_abs,
    xl_rowcol_to_cell as upstream_cell,
)

import mojo_xlsxwriter as ours
from mojo_xlsxwriter.exceptions import (
    DuplicateWorksheetName,
    InvalidWorksheetName,
    OverlappingRange,
)
from mojo_xlsxwriter.utility import (
    xl_col_to_name,
    xl_range,
    xl_range_abs,
    xl_rowcol_to_cell,
)

MAIN = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def make_bytes(workbook_class, populate, options=None):
    destination = BytesIO()
    workbook = workbook_class(destination, {"in_memory": True, **(options or {})})
    populate(workbook)
    workbook.close()
    return destination.getvalue()


def load(data, data_only=False):
    return openpyxl.load_workbook(BytesIO(data), data_only=data_only)


def shared_xml(data):
    with zipfile.ZipFile(BytesIO(data)) as archive:
        return archive.read("xl/sharedStrings.xml")


def worksheet_xml(data, index=1):
    with zipfile.ZipFile(BytesIO(data)) as archive:
        return ET.fromstring(archive.read(f"xl/worksheets/sheet{index}.xml"))


def test_shared_strings_xml_is_byte_identical_to_upstream():
    values = [
        "repeat",
        "repeat",
        " leading",
        "trailing ",
        "a&<b\x01_x0000_",
        "β中 café",
        "\ufffe\uffff",
        "",
    ]

    def populate(workbook):
        sheet = workbook.add_worksheet()
        for row, value in enumerate(values):
            sheet.write_string(row, 0, value)

    assert shared_xml(make_bytes(ours.Workbook, populate)) == shared_xml(
        make_bytes(xlsxwriter.Workbook, populate)
    )


def test_shared_string_order_and_count_include_overwrites():
    def populate(workbook):
        sheet = workbook.add_worksheet()
        sheet.write(0, 0, "first")
        sheet.write(0, 0, "second")
        sheet.write(0, 1, "second")

    ours_xml = ET.fromstring(shared_xml(make_bytes(ours.Workbook, populate)))
    theirs_xml = ET.fromstring(shared_xml(make_bytes(xlsxwriter.Workbook, populate)))
    assert ours_xml.attrib == theirs_xml.attrib == {"count": "3", "uniqueCount": "2"}
    ours_text = [node.text for node in ours_xml.findall("m:si/m:t", MAIN)]
    theirs_text = [node.text for node in theirs_xml.findall("m:si/m:t", MAIN)]
    assert ours_text == theirs_text == ["first", "second"]
    assert load(make_bytes(ours.Workbook, populate)).active["A1"].value == "second"


def comprehensive(workbook):
    data = workbook.add_worksheet("Data & More")
    other = workbook.add_worksheet("Other")
    header = workbook.add_format(
        {
            "bold": True,
            "font_color": "white",
            "bg_color": "#4472C4",
            "border": 1,
            "align": "center",
        }
    )
    money = workbook.add_format({"num_format": "$#,##0.00", "border": 1})
    date_format = workbook.add_format({"num_format": "yyyy-mm-dd hh:mm"})
    data.write_row("A1", ["Name", "Value", "Active", "Formula", "Date"], header)
    data.write_row(1, 0, ["alpha", 12.5, True])
    data.write_formula(1, 3, "=B2*2", money, 25)
    data.write_datetime(1, 4, dt.datetime(2024, 2, 29, 13, 45), date_format)
    data.write_row(2, 0, ["beta", -3, False])
    data.write_formula(2, 3, "=SUM(B2:B3)", money, 9.5)
    data.write_datetime(2, 4, dt.date(2020, 1, 1), date_format)
    data.write("A4", " spaced ")
    data.write("B4", "β & <xml>")
    data.write_blank("C4", None, header)
    data.merge_range("A6:C6", "Merged title", header)
    data.write_url("A8", "https://example.com/path?a=1&b=2", None, "Example", "tip")
    data.write_url("B8", "internal:Other!A1", None, "Jump")
    data.set_column("A:A", 18)
    data.set_column(1, 4, 14, money)
    data.set_row(0, 24)
    data.freeze_panes(1, 1)
    data.autofilter("A1:E4")
    other.write("A1", "target")
    workbook.set_properties({"author": "Parity", "title": "Workbook"})


def test_comprehensive_workbook_behavior_matches_upstream():
    ours_data = make_bytes(ours.Workbook, comprehensive)
    theirs_data = make_bytes(xlsxwriter.Workbook, comprehensive)
    ours_book = load(ours_data)
    theirs_book = load(theirs_data)
    assert ours_book.sheetnames == theirs_book.sheetnames
    for sheet_name in ours_book.sheetnames:
        ours_sheet = ours_book[sheet_name]
        theirs_sheet = theirs_book[sheet_name]
        for row in range(1, max(ours_sheet.max_row, theirs_sheet.max_row) + 1):
            for col in range(1, max(ours_sheet.max_column, theirs_sheet.max_column) + 1):
                ours_cell = ours_sheet.cell(row, col)
                theirs_cell = theirs_sheet.cell(row, col)
                assert ours_cell.value == theirs_cell.value
                assert ours_cell.data_type == theirs_cell.data_type
                assert ours_cell.number_format == theirs_cell.number_format
        assert list(ours_sheet.merged_cells.ranges) == list(theirs_sheet.merged_cells.ranges)
    ours_sheet = ours_book["Data & More"]
    theirs_sheet = theirs_book["Data & More"]
    assert ours_sheet.freeze_panes == theirs_sheet.freeze_panes == "B2"
    assert ours_sheet.auto_filter.ref == theirs_sheet.auto_filter.ref == "A1:E4"
    assert ours_sheet.column_dimensions["A"].width == pytest.approx(
        theirs_sheet.column_dimensions["A"].width
    )
    assert ours_sheet.row_dimensions[1].height == theirs_sheet.row_dimensions[1].height
    assert ours_sheet["A1"].font.bold == theirs_sheet["A1"].font.bold
    assert ours_sheet["A1"].font.color.rgb == theirs_sheet["A1"].font.color.rgb
    assert ours_sheet["A1"].fill.fgColor.rgb == theirs_sheet["A1"].fill.fgColor.rgb
    assert ours_sheet["A8"].hyperlink.target == theirs_sheet["A8"].hyperlink.target
    assert ours_sheet["A8"].hyperlink.tooltip == theirs_sheet["A8"].hyperlink.tooltip
    assert ours_sheet["B8"].hyperlink.location == theirs_sheet["B8"].hyperlink.location


def test_formula_cached_values_match_upstream_xml():
    def populate(workbook):
        sheet = workbook.add_worksheet()
        sheet.write_formula(0, 0, "=1+2", None, 3)
        sheet.write_formula(0, 1, "=TRUE()", None, True)
        sheet.write_formula(0, 2, '="ok"', None, "ok")
        sheet.write_formula(0, 3, "=1/0", None, "#DIV/0!")

    ours_root = worksheet_xml(make_bytes(ours.Workbook, populate))
    theirs_root = worksheet_xml(make_bytes(xlsxwriter.Workbook, populate))
    for reference in ("A1", "B1", "C1", "D1"):
        query = f".//m:c[@r='{reference}']"
        ours_cell = ours_root.find(query, MAIN)
        theirs_cell = theirs_root.find(query, MAIN)
        assert ours_cell.attrib.get("t") == theirs_cell.attrib.get("t")
        assert ours_cell.find("m:f", MAIN).text == theirs_cell.find("m:f", MAIN).text
        assert ours_cell.find("m:v", MAIN).text == theirs_cell.find("m:v", MAIN).text


@pytest.mark.parametrize("date_1904", [False, True])
def test_datetime_serials_and_default_format_match_upstream(date_1904):
    values = [
        dt.datetime(1900, 1, 1, 12),
        dt.datetime(1900, 2, 28),
        dt.datetime(1900, 3, 1),
        dt.date(2024, 1, 2),
        dt.time(6, 30),
        dt.timedelta(days=2, seconds=1),
    ]
    options = {"date_1904": date_1904, "default_date_format": "yyyy-mm-dd hh:mm:ss"}

    def populate(workbook):
        sheet = workbook.add_worksheet()
        for row, value in enumerate(values):
            sheet.write_datetime(row, 0, value)

    ours_root = worksheet_xml(make_bytes(ours.Workbook, populate, options))
    theirs_root = worksheet_xml(make_bytes(xlsxwriter.Workbook, populate, options))
    ours_values = [node.text for node in ours_root.findall(".//m:c/m:v", MAIN)]
    theirs_values = [node.text for node in theirs_root.findall(".//m:c/m:v", MAIN)]
    assert [float(value) for value in ours_values] == pytest.approx(
        [float(value) for value in theirs_values]
    )
    book = load(make_bytes(ours.Workbook, populate, options))
    assert all(
        book.active.cell(row, 1).number_format == "yyyy-mm-dd hh:mm:ss"
        for row in range(1, len(values) + 1)
    )


def test_write_dispatch_options_match_upstream():
    options = {
        "strings_to_numbers": True,
        "strings_to_formulas": False,
        "strings_to_urls": False,
        "nan_inf_to_errors": True,
    }

    def populate(workbook):
        sheet = workbook.add_worksheet()
        for col, value in enumerate(["12.5", "=1+2", "https://example.com", math.nan, math.inf]):
            sheet.write(0, col, value)

    ours_root = worksheet_xml(make_bytes(ours.Workbook, populate, options))
    theirs_root = worksheet_xml(make_bytes(xlsxwriter.Workbook, populate, options))
    ours_cells = ours_root.findall(".//m:c", MAIN)
    theirs_cells = theirs_root.findall(".//m:c", MAIN)
    assert [cell.attrib.get("t") for cell in ours_cells] == [
        cell.attrib.get("t") for cell in theirs_cells
    ]
    assert load(make_bytes(ours.Workbook, populate, options)).active["A1"].value == 12.5


def test_write_row_preserves_string_dispatch():
    options = {"strings_to_numbers": True}

    def populate(workbook):
        sheet = workbook.add_worksheet()
        sheet.write_row(0, 0, ["12.5", "=1+2", "https://example.com"])

    ours_data = make_bytes(ours.Workbook, populate, options)
    theirs_data = make_bytes(xlsxwriter.Workbook, populate, options)
    ours_sheet = load(ours_data, data_only=False).active
    theirs_sheet = load(theirs_data, data_only=False).active
    assert [cell.value for cell in ours_sheet[1]] == [
        cell.value for cell in theirs_sheet[1]
    ]
    assert ours_sheet["C1"].hyperlink.target == theirs_sheet["C1"].hyperlink.target


def test_nan_inf_raise_without_option():
    for workbook_class in (ours.Workbook, xlsxwriter.Workbook):
        destination = BytesIO()
        workbook = workbook_class(destination, {"in_memory": True})
        sheet = workbook.add_worksheet()
        with pytest.raises(TypeError, match="NAN/INF"):
            sheet.write_number(0, 0, math.nan)


def test_a1_and_numeric_call_forms_are_equivalent():
    def populate(workbook):
        sheet = workbook.add_worksheet()
        cell_format = workbook.add_format({"bold": True})
        assert sheet.write_string("A1", "a", cell_format) == 0
        assert sheet.write_number("B1", 2.5, cell_format) == 0
        assert sheet.write_boolean("C1", True, cell_format) == 0
        assert sheet.write_formula("D1", "=1+2", cell_format, 3) == 0
        assert sheet.write_row("A2", [1, 2], cell_format) == 0
        assert sheet.write_column("C2", [3, 4], cell_format) == 0
        sheet.set_column("A:C", 20, cell_format)
        sheet.merge_range("A5:C5", "merged", cell_format)

    ours_book = load(make_bytes(ours.Workbook, populate))
    theirs_book = load(make_bytes(xlsxwriter.Workbook, populate))
    for row in ours_book.active.iter_rows():
        assert [cell.value for cell in row] == [
            theirs_book.active.cell(cell.row, cell.column).value for cell in row
        ]


def test_return_codes_and_bounds_match_upstream():
    long_string = "x" * 40_000
    for workbook_class in (ours.Workbook, xlsxwriter.Workbook):
        workbook = workbook_class(BytesIO(), {"in_memory": True})
        sheet = workbook.add_worksheet()
        assert sheet.write_string(0, 0, long_string) == -2
        assert sheet.write_string(1_048_576, 0, "bad") == -1
        assert sheet.write_number(0, 16_384, 1) == -1
        assert sheet.set_row(-1, 20) == -1
        assert sheet.set_column(-1, 0, 20) == -1
        workbook.close()


def test_empty_workbook_has_no_shared_string_part_like_upstream():
    def populate(workbook):
        workbook.add_worksheet()

    for workbook_class in (ours.Workbook, xlsxwriter.Workbook):
        with zipfile.ZipFile(BytesIO(make_bytes(workbook_class, populate))) as archive:
            assert "xl/sharedStrings.xml" not in archive.namelist()


def test_constant_memory_is_explicitly_outside_shared_string_subset():
    with pytest.raises(NotImplementedError, match="inline strings"):
        ours.Workbook(BytesIO(), {"constant_memory": True})


def test_sheet_name_validation_and_lookup():
    workbook = ours.Workbook(BytesIO(), {"in_memory": True})
    sheet = workbook.add_worksheet("Data")
    assert workbook.get_worksheet_by_name("data") is sheet
    assert workbook.worksheets() == [sheet]
    with pytest.raises(DuplicateWorksheetName):
        workbook.add_worksheet("DATA")
    with pytest.raises(InvalidWorksheetName):
        workbook.add_worksheet("bad/name")
    with pytest.raises(InvalidWorksheetName):
        workbook.add_worksheet("x" * 32)
    workbook.close()


def test_merge_overlap_is_rejected():
    workbook = ours.Workbook(BytesIO(), {"in_memory": True})
    sheet = workbook.add_worksheet()
    sheet.merge_range("A1:C2", "first")
    with pytest.raises(OverlappingRange):
        sheet.merge_range("C2:D3", "second")
    workbook.close()


def test_utility_functions_match_upstream():
    for col in (0, 25, 26, 701, 702, 16_383):
        assert xl_col_to_name(col) == upstream_col_name(col)
        assert xl_col_to_name(col, True) == upstream_col_name(col, True)
    for row, col in ((0, 0), (10, 27), (1_048_575, 16_383)):
        assert xl_rowcol_to_cell(row, col) == upstream_cell(row, col)
        assert xl_rowcol_to_cell(row, col, True, True) == upstream_cell(
            row, col, True, True
        )
    assert xl_range(0, 0, 9, 4) == upstream_range(0, 0, 9, 4)
    assert xl_range_abs(0, 0, 9, 4) == upstream_range_abs(0, 0, 9, 4)


def test_documented_format_properties_match_upstream():
    properties = {
        "bold": True,
        "italic": True,
        "underline": 2,
        "font_name": "Arial",
        "font_size": 14,
        "font_color": "#123456",
        "bg_color": "#ABCDEF",
        "num_format": "0.000",
        "align": "center",
        "valign": "top",
        "text_wrap": True,
        "indent": 2,
        "rotation": 30,
        "border": 2,
        "border_color": "#654321",
        "locked": False,
        "hidden": True,
    }

    def populate(workbook):
        sheet = workbook.add_worksheet()
        sheet.write(0, 0, 1.25, workbook.add_format(properties))

    ours_cell = load(make_bytes(ours.Workbook, populate)).active["A1"]
    theirs_cell = load(make_bytes(xlsxwriter.Workbook, populate)).active["A1"]
    assert vars(ours_cell.font) == vars(theirs_cell.font)
    assert vars(ours_cell.fill) == vars(theirs_cell.fill)
    assert vars(ours_cell.border) == vars(theirs_cell.border)
    assert vars(ours_cell.alignment) == vars(theirs_cell.alignment)
    assert vars(ours_cell.protection) == vars(theirs_cell.protection)
    assert ours_cell.number_format == theirs_cell.number_format


def test_documented_context_path_timezone_and_zip64_surface(tmp_path):
    destination = tmp_path / "surface.xlsx"
    with ours.Workbook(
        destination,
        {"remove_timezone": True, "use_zip64": True},
    ) as workbook:
        assert workbook.allow_zip64
        sheet = workbook.add_worksheet("Surface")
        sheet.write_datetime(
            "A1",
            dt.datetime(2026, 7, 30, tzinfo=dt.timezone.utc),
        )
    assert openpyxl.load_workbook(destination).active["A1"].value == pytest.approx(
        46233
    )
